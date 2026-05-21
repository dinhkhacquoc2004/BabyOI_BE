from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
CHANGELOG_DIR = ROOT / "src" / "main" / "resources" / "db" / "changelog" / "changes"
OUTPUT = ROOT / "food-image-checklist.docx"


def fix_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value)
    # Some seed files contain Vietnamese text that has already been mojibake-encoded.
    # Recover it for the human checklist when possible.
    if any(mark in text for mark in ("Ã", "á»", "Ä", "Æ")):
        for encoding in ("cp1252", "latin1"):
            try:
                fixed = text.encode(encoding).decode("utf-8")
                if fixed.count("�") < text.count("�"):
                    return fixed
                if any(ch in fixed for ch in "ăâđêôơưĂÂĐÊÔƠƯ"):
                    return fixed
            except UnicodeError:
                continue
    return text


def parse_sql_value(raw: str) -> object:
    raw = raw.strip()
    if raw.upper() == "NULL":
        return None
    if raw.startswith("'") and raw.endswith("'"):
        return fix_text(raw[1:-1].replace("''", "'"))
    try:
        if "." in raw:
            return float(raw)
        return int(raw)
    except ValueError:
        return fix_text(raw)


def split_tuple(content: str) -> list[object]:
    values: list[str] = []
    buf: list[str] = []
    in_quote = False
    i = 0
    while i < len(content):
        ch = content[i]
        if ch == "'":
            buf.append(ch)
            if in_quote and i + 1 < len(content) and content[i + 1] == "'":
                buf.append("'")
                i += 2
                continue
            in_quote = not in_quote
        elif ch == "," and not in_quote:
            values.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
        i += 1
    if buf:
        values.append("".join(buf))
    return [parse_sql_value(value) for value in values]


def tuple_at(text: str, start: int) -> tuple[str, int] | None:
    depth = 0
    in_quote = False
    i = start
    while i < len(text):
        ch = text[i]
        if ch == "'":
            if in_quote and i + 1 < len(text) and text[i + 1] == "'":
                i += 2
                continue
            in_quote = not in_quote
        elif not in_quote:
            if ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
                if depth == 0:
                    return text[start + 1 : i], i + 1
        i += 1
    return None


def parse_food_rows() -> list[dict[str, object]]:
    rows: dict[int, dict[str, object]] = {}
    for file_name in ("003_seed_food_catalog.sql", "004_add_food_goal_catalog.sql"):
        text = (CHANGELOG_DIR / file_name).read_text(encoding="utf-8")
        for match in re.finditer(r"(?m)^\s*\(\s*\d+\s*,", text):
            tuple_start = text.find("(", match.start(), match.end())
            parsed = tuple_at(text, tuple_start)
            if not parsed:
                continue
            values = split_tuple(parsed[0])
            if len(values) < 10:
                continue
            if not isinstance(values[0], int) or values[1] not in (1, 2) or not isinstance(values[2], str):
                continue

            food_id = int(values[0])
            if len(values) >= 18 and isinstance(values[3], str):
                row = {
                    "id": food_id,
                    "function_code": values[1],
                    "name": values[2],
                    "advance_for": values[3],
                    "protein": values[4],
                    "carb": values[5],
                    "veg": values[6],
                    "calories": values[10],
                }
            else:
                row = {
                    "id": food_id,
                    "function_code": values[1],
                    "name": values[2],
                    "advance_for": "",
                    "protein": values[3],
                    "carb": values[4],
                    "veg": values[5],
                    "calories": values[6],
                }
            rows.setdefault(food_id, row)
    return [rows[key] for key in sorted(rows)]


def parse_ingredient_rows() -> list[dict[str, object]]:
    text = (CHANGELOG_DIR / "003_seed_food_catalog.sql").read_text(encoding="utf-8")
    match = re.search(r"FROM\s+unnest\(ARRAY\[(.*?)\]::text\[\]\)", text, flags=re.S | re.I)
    if not match:
        return []
    names = [fix_text(item.replace("''", "'")) for item in re.findall(r"'((?:''|[^'])*)'", match.group(1))]
    return [{"id": index, "name": name} for index, name in enumerate(names, start=1)]


def target_label(function_code: object) -> str:
    if function_code == 1:
        return "Mẹ"
    if function_code == 2:
        return "Bé"
    return ""


def add_table(document: Document, headers: list[str], rows: list[list[object]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        header_cells[index].text = header
        for paragraph in header_cells[index].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = fix_text(value)


def main() -> None:
    foods = parse_food_rows()
    ingredients = parse_ingredient_rows()

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(8)

    document.add_heading("Checklist ảnh cho Food Library và Ingredients", level=1)
    document.add_paragraph(
        f"Tổng món ăn: {len(foods)}. Tổng nguyên liệu: {len(ingredients)}. "
        "Điền link ảnh vào cột Image URL / Ghi chú ảnh rồi có thể nhập lại vào DB."
    )

    document.add_heading("1. Danh sách món ăn", level=2)
    add_table(
        document,
        ["Food ID", "Tên món ăn", "Dành cho", "Advice For", "Protein", "Carb", "Rau/củ/trái cây", "Calo", "Image URL / Ghi chú ảnh"],
        [
            [
                food["id"],
                food["name"],
                target_label(food["function_code"]),
                food["advance_for"],
                food["protein"],
                food["carb"],
                food["veg"],
                food["calories"],
                "",
            ]
            for food in foods
        ],
    )

    document.add_page_break()
    document.add_heading("2. Danh sách nguyên liệu", level=2)
    add_table(
        document,
        ["Ingredient ID", "Tên nguyên liệu", "Image URL / Ghi chú ảnh"],
        [[ingredient["id"], ingredient["name"], ""] for ingredient in ingredients],
    )

    document.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"foods={len(foods)} ingredients={len(ingredients)}")


if __name__ == "__main__":
    main()
